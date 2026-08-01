from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TARA_ROOT = Path("/Users/vigneshramoo/Documents/TARA")
OUTPUT_DIR = TARA_ROOT / "TARA & Fiends"
HANDBOOK_DOCX = OUTPUT_DIR / "TARA Crew Scent Handbook - AUREYA ZEPHYR MARIS ELIORA - 2026-05-14.docx"
LOGO_PATH = TARA_ROOT / "Marketing & Branding/tara-logo/PNG/tara-workmark_Main wordmark-.png"

SOURCE_FILES = {
    "AUREYA": TARA_ROOT / "Marketing & Branding/AUREYA Marketing Notes/TARA_AUREYA_MarketingAccordsAndNotes_FINAL_2026-05-07.md",
    "ZEPHYR": TARA_ROOT / "Marketing & Branding/ZEPHYR Marketing Notes/TARA_ZEPHYR_MarketingAccordsAndNotes_FINAL_2026-05-07.md",
    "MARIS": TARA_ROOT / "Marketing & Branding/MARIS Marketing Notes/TARA_MARIS_MarketingAccordsAndNotes_FINAL_2026-05-07.md",
    "ELIORA": TARA_ROOT / "Marketing & Branding/ELIORA Marketing Notes/TARA_ELIORA_MarketingAccordsAndNotes_FINAL_2026-05-07.md",
}

BLACK = RGBColor(10, 10, 10)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(84, 84, 84)
GOLD = RGBColor(202, 158, 91)
LIGHT_GOLD = "F7F1E7"
LINE = "D8C7AD"


@dataclass
class ScentNotes:
    name: str
    positioning: str
    public_positioning: str
    top_accords: list[list[str]]
    pyramid: list[list[str]]
    short_description: str
    product_description: str
    launch_description: str
    note_language: list[str]
    primary_hook: str
    secondary_hooks: list[str]
    script_heading: str
    script_cues: dict[str, str]
    guardrails: list[str]
    final_direction: str


def section(markdown: str, heading: str) -> str:
    pattern = rf"^## {re.escape(heading)}\s*\n(.*?)(?=^## |\Z)"
    match = re.search(pattern, markdown, flags=re.M | re.S)
    return match.group(1).strip() if match else ""


def clean_markdown_text(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = text.replace("  \n", "\n")
    return text.strip()


def parse_table(section_text: str) -> list[list[str]]:
    rows: list[list[str]] = []

    for line in section_text.splitlines():
        line = line.strip()
        if not line.startswith("|") or re.fullmatch(r"\|[\s:|\-]+\|", line):
            continue

        cells = [clean_markdown_text(cell.strip()) for cell in line.strip("|").split("|")]
        if cells:
            rows.append(cells)

    return rows


def parse_bullets(section_text: str) -> list[str]:
    return [
        clean_markdown_text(match.group(1))
        for match in re.finditer(r"^- (.+)$", section_text, flags=re.M)
    ]


def parse_blockquote(section_text: str) -> str:
    match = re.search(r"^> (.+)$", section_text, flags=re.M)
    return clean_markdown_text(match.group(1)) if match else ""


def parse_description(section_text: str, label: str) -> str:
    pattern = rf"{re.escape(label)}:\s*\n\n(.*?)(?=\n\n[A-Z][A-Za-z ]+ description:|\Z)"
    match = re.search(pattern, section_text, flags=re.S)
    return clean_markdown_text(match.group(1)) if match else ""


def parse_script_cues(section_text: str) -> dict[str, str]:
    cues = {}

    for label in ["Opening cue", "Heart cue", "Drydown cue"]:
        match = re.search(rf"{label}:\s*\n\n\"(.*?)\"", section_text, flags=re.S)
        if match:
            cues[label.replace(" cue", "")] = clean_markdown_text(match.group(1))

    return cues


def parse_notes(name: str, path: Path) -> ScentNotes:
    markdown = path.read_text()
    positioning_text = section(markdown, "Positioning Summary")
    approved_text = section(markdown, "Approved Product Descriptions")
    story_text = section(markdown, "Storytelling Hooks")
    sales_script = section(markdown, "Sales Script Note Cues")
    booth_script = section(markdown, "Booth Script Note Cues")
    script_text = sales_script or booth_script

    paragraphs_before_public = positioning_text.split("Public positioning:")[0].strip()
    final_direction = clean_markdown_text(section(markdown, "Final Marketing Direction"))

    return ScentNotes(
        name=name,
        positioning=clean_markdown_text(paragraphs_before_public),
        public_positioning=parse_blockquote(positioning_text),
        top_accords=parse_table(section(markdown, "Top Accords")),
        pyramid=parse_table(section(markdown, "Marketing Note Pyramid")),
        short_description=parse_description(approved_text, "Short description"),
        product_description=parse_description(approved_text, "Product page description"),
        launch_description=parse_description(approved_text, "Launch description"),
        note_language=parse_bullets(section(markdown, "Suggested Note Language")),
        primary_hook=parse_blockquote(story_text),
        secondary_hooks=parse_bullets(story_text),
        script_heading="Sales Script" if sales_script else "Booth Script",
        script_cues=parse_script_cues(script_text),
        guardrails=parse_bullets(section(markdown, "Marketing Guardrails")),
        final_direction=final_direction,
    )


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))

    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)

    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")

    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ["top", "left", "bottom", "right"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")

    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)

    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run(run, *, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def paragraph(doc: Document, text: str = "", *, style: str | None = None, before: int = 0, after: int = 6, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run(run, size=11, color=INK)
    return p


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run, size={1: 17, 2: 13, 3: 11.5}[level], color=GOLD if level <= 2 else BLACK, bold=True)
    return p


def bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    set_run(run, size=10.5, color=INK)
    return p


def add_callout(doc: Document, text: str, label: str | None = None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_fill(cell, LIGHT_GOLD)
    set_cell_border(cell)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if label:
        label_run = p.add_run(f"{label.upper()}: ")
        set_run(label_run, size=9, color=GOLD, bold=True)
    body_run = p.add_run(text)
    set_run(body_run, size=10.5, color=INK)
    paragraph(doc, "", after=4)


def add_table(doc: Document, rows: list[list[str]], widths: list[float], header: bool = True):
    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for col_index, width in enumerate(widths):
        table.columns[col_index].width = Inches(width)

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if header and row_index == 0:
                set_cell_fill(cell, "111111")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run(
                run,
                size=9.5 if not (header and row_index == 0) else 9,
                color=RGBColor(255, 255, 255) if header and row_index == 0 else INK,
                bold=header and row_index == 0,
            )

    paragraph(doc, "", after=4)


def add_section_header(section, title: str):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    run.add_picture(str(LOGO_PATH), width=Inches(1.15))
    p.add_run("   ")
    title_run = p.add_run(title)
    set_run(title_run, size=8.5, color=MUTED, bold=True)


def add_cover(doc: Document):
    section = doc.sections[0]
    add_section_header(section, "Crew Scent Handbook")

    paragraph(doc, "", after=20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO_PATH), width=Inches(4.2))
    paragraph(doc, "", after=16)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Crew Scent Handbook")
    set_run(title_run, size=26, color=BLACK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle_run = subtitle.add_run("AUREYA / ZEPHYR / MARIS / ELIORA - Special Launch")
    set_run(subtitle_run, size=13, color=GOLD, bold=True)

    add_callout(
        doc,
        "A fast, polished guide for TARA crews to explain each scent clearly, protect the brand language, and help customers find the fragrance that feels like them.",
        "Purpose",
    )
    paragraph(
        doc,
        f"Prepared for TARA & Friends crews | {datetime.now().strftime('%d %B %Y')}",
        after=18,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def add_quick_start(doc: Document, notes: list[ScentNotes]):
    doc.add_page_break()
    heading(doc, "Crew Quick Start", 1)
    add_callout(
        doc,
        "Lead with mood, not ingredients. Ask what the customer wants to feel, then use TARA's approved scent language to guide the test.",
        "Booth rhythm",
    )

    heading(doc, "Social Handles To Follow", 2)
    for item in [
        "Instagram: @tara_scents.my",
        "Website: tarascents.com",
        "Link hub: tarascents.com/links",
        "WhatsApp Concierge: +60 11-4304 2883",
    ]:
        bullet(doc, item)

    heading(doc, "One-Glance Scent Map", 2)
    rows = [["Scent", "Position", "Primary Hook", "Customer Direction"]]
    for note in notes:
        direction = {
            "AUREYA": "Soft, luminous, feminine warmth.",
            "ZEPHYR": "Clean masculine freshness with warmth.",
            "MARIS": "Unisex mineral marine woods.",
            "ELIORA": "Limited golden floral-musk evening secret.",
        }[note.name]
        rows.append([note.name, direction, note.primary_hook, note.short_description])
    add_table(doc, rows, [0.9, 1.55, 1.65, 2.4])

    heading(doc, "Crew Guardrails", 2)
    for item in [
        "Never say clone, dupe, mix, or inspired-by.",
        "Do not mention internal formula construction, ratios, source blends, or technical polish levels.",
        "Use note language as olfactive impression, not ingredient disclosure.",
        "For ELIORA, keep the story limited, mysterious, and pop-up exclusive.",
        "Keep the sales tone warm, helpful, and quietly confident.",
    ]:
        bullet(doc, item)


def add_scent(doc: Document, note: ScentNotes):
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_section_header(doc.sections[-1], f"{note.name} | TARA Crew Scent Handbook")
    heading(doc, note.name if note.name != "ELIORA" else "ELIORA - Special Launch", 1)
    if note.primary_hook:
        add_callout(doc, note.primary_hook, "Primary hook")

    heading(doc, "Positioning", 2)
    paragraph(doc, note.positioning)
    add_callout(doc, note.public_positioning, "Public line")

    heading(doc, note.script_heading, 2)
    script_rows = [["Moment", "What to say"]]
    for cue_name in ["Opening", "Heart", "Drydown"]:
        if cue_name in note.script_cues:
            script_rows.append([cue_name, note.script_cues[cue_name]])
    add_table(doc, script_rows, [1.15, 5.35])

    heading(doc, "Marketing Note Pyramid", 2)
    add_table(doc, note.pyramid, [1.15, 5.35])

    heading(doc, "Top Accord Roles", 2)
    add_table(doc, note.top_accords, [1.85, 4.65])

    heading(doc, "Approved Product Copy", 2)
    for label, copy in [
        ("Short", note.short_description),
        ("Product page", note.product_description),
        ("Launch", note.launch_description),
    ]:
        add_callout(doc, copy, label)

    heading(doc, "Language To Reuse", 2)
    for phrase in note.note_language:
        bullet(doc, phrase)

    heading(doc, "Guardrails", 2)
    for guardrail in note.guardrails:
        bullet(doc, guardrail)

    heading(doc, "Final Direction", 2)
    paragraph(doc, note.final_direction)


def build_document():
    notes = [parse_notes(name, SOURCE_FILES[name]) for name in ["AUREYA", "ZEPHYR", "MARIS", "ELIORA"]]
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.76)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"

    add_cover(doc)
    add_quick_start(doc, notes)
    for note in notes:
        add_scent(doc, note)

    section = doc.sections[-1]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("TARA Crew Scent Handbook | Internal Guide")
    set_run(run, size=8.5, color=MUTED)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(HANDBOOK_DOCX)
    return HANDBOOK_DOCX


if __name__ == "__main__":
    print(build_document())
