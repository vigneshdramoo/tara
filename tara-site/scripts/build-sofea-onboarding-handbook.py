from __future__ import annotations

import importlib.util
import sys
from dataclasses import dataclass
from datetime import date
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


TARA_ROOT = Path("/Users/vigneshramoo/Documents/TARA")
SITE_DIR = TARA_ROOT / "tara-site"
OUTPUT_DIR = TARA_ROOT / "TARA & Fiends"
HANDBOOK_DIR = OUTPUT_DIR / "Handbooks"
PREPARED_DATE = date(2026, 6, 12)
DATE_STAMP = PREPARED_DATE.strftime("%Y-%m-%d")
DISPLAY_DATE = PREPARED_DATE.strftime("%d %B %Y")

CREW_FIRST_NAME = "Sofea"
CREW_FULL_NAME = "Nur Zetty Sofea Binti Abdul Rashid"
CREW_ROLE = "Part-Time Perfume Promoter"
HANDBOOK_TITLE = f"{CREW_FIRST_NAME} TARA Onboarding Handbook"
HANDBOOK_SUBTITLE = "TARA Scent Trail | Scent Education + Stop 03 Readiness"
EVENT_NAME = "TARA Scent Trail: Stop 03"
EVENT_VENUE = "KL Sidewalk @ Huuha Land"
EVENT_DATES = "12 June - 14 June 2026"
EVENT_HOURS = "5:00PM - 12:00AM daily"
EVENT_STATUS = "Upcoming"

DOCX_OUT = HANDBOOK_DIR / f"{HANDBOOK_TITLE} - {DATE_STAMP}.docx"
PDF_OUT = HANDBOOK_DIR / f"{HANDBOOK_TITLE} - {DATE_STAMP}.pdf"

LOGO_DOCX = TARA_ROOT / "Marketing & Branding/tara-logo/PNG/tara-workmark_Main wordmark-.png"
LOGO_PDF = SITE_DIR / "netlify/functions/lib/assets/tara-official-wordmark.png"
BRAND_BOOK = TARA_ROOT / "Marketing & Branding/Storytelling/TARA/TARA_CoreTrio_BrandBook.pdf"
CREW_BUILDER = SITE_DIR / "scripts/build-crew-scent-handbook.py"

BLACK = RGBColor(10, 10, 10)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(84, 84, 84)
GOLD = RGBColor(202, 158, 91)
WHITE = RGBColor(255, 255, 255)
LIGHT_GOLD = "F7F1E7"
LIGHT_GRAY = "F7F7F5"
LINE = "D8C7AD"

PDF_BLACK = colors.HexColor("#0A0A0A")
PDF_INK = colors.HexColor("#202020")
PDF_MUTED = colors.HexColor("#565656")
PDF_GOLD = colors.HexColor("#CA9E5B")
PDF_LIGHT_GOLD = colors.HexColor("#F7F1E7")
PDF_LIGHT_GRAY = colors.HexColor("#F7F7F5")
PDF_LINE = colors.HexColor("#D8C7AD")


spec = importlib.util.spec_from_file_location("crew_handbook_builder", CREW_BUILDER)
builder = importlib.util.module_from_spec(spec)
assert spec and spec.loader
sys.modules["crew_handbook_builder"] = builder
spec.loader.exec_module(builder)


@dataclass(frozen=True)
class ScentTraining:
    identity: str
    customer_fit: list[str]
    story_anchor: str
    thirty_second_pitch: str
    try_when: str
    trainer_note: str


SCENT_TRAINING = {
    "AUREYA": ScentTraining(
        identity="Radiant identity, soft confidence, and golden memory.",
        customer_fit=[
            "Wants a graceful feminine scent that feels polished, warm, and emotionally soft.",
            "Likes florals that are wearable instead of loud or overly sweet.",
            "Needs a first-impression scent for daytime, dates, gatherings, or gifting.",
        ],
        story_anchor=(
            "AUREYA is the sunrise scent: pear brightness, white petals, soft musk, and amber warmth. "
            "The feeling is gentle confidence - someone who glows without forcing attention."
        ),
        thirty_second_pitch=(
            "AUREYA starts bright and elegant, then becomes soft and floral before settling into a warm "
            "musk-amber glow. It is feminine, polished, and easy to love."
        ),
        try_when="When someone asks for pretty, soft, elegant, feminine, warm, or easy to wear.",
        trainer_note="Keep the language tender and luminous. Think morning light, silk, petals, and warmth.",
    ),
    "ZEPHYR": ScentTraining(
        identity="Controlled presence, bright air, tailored warmth.",
        customer_fit=[
            "Wants a clean masculine scent that feels modern and confident.",
            "Likes fresh openings but still wants warmth and a memorable trail.",
            "Needs a daily scent for work, campus, meetings, or smart casual wear.",
        ],
        story_anchor=(
            "ZEPHYR is the composed one: sparkling citrus and clean air moving into polished woods, musks, "
            "and ambroxan warmth. It draws compliments without shouting."
        ),
        thirty_second_pitch=(
            "ZEPHYR opens clean and bright, then turns woody and tailored on skin. It feels fresh, confident, "
            "and quietly magnetic - the kind of scent people notice up close."
        ),
        try_when="When someone asks for fresh, clean, masculine, professional, compliment-getting, or daily wear.",
        trainer_note="Use confident words. Keep it crisp, clean, and composed rather than sporty or aggressive.",
    ),
    "MARIS": ScentTraining(
        identity="Calm confidence, salt air, mineral light, skin-close woods.",
        customer_fit=[
            "Wants something unisex, clean, quiet, and a little mysterious.",
            "Likes marine freshness but does not want a generic beach scent.",
            "Needs a subtle signature for daily wear, travel, or close moments.",
        ],
        story_anchor=(
            "MARIS is not sunny beach freshness. It is midnight coastline energy: salt air, green sage, "
            "a mineral amber beam, and sandalwood-musk calm."
        ),
        thirty_second_pitch=(
            "MARIS starts with cool salt air and sage, warms into mineral amber, then dries down to sandalwood, "
            "driftwood, and skin musk. It is clean, unisex, and quietly magnetic."
        ),
        try_when="When someone asks for unisex, calm, clean, subtle, oceanic, mineral, or not-too-sweet.",
        trainer_note="Avoid beach-holiday language. Say salt, signal, silence, fog, wood, and skin.",
    ),
    "ELIORA": ScentTraining(
        identity="Pop-up-only golden floral musk, mystery, intimacy, and limited discovery.",
        customer_fit=[
            "Wants something more exclusive, feminine, warm, and evening-coded.",
            "Likes florals with depth, musk, amber, and a soft sensual trail.",
            "Responds to limited launches, story, and the feeling of discovering a secret.",
        ],
        story_anchor=(
            "ELIORA is the secret scent. It is unveiled where TARA appears: golden shimmer, creamy white florals, "
            "soft spice, clean musk, amber skin, and velvet warmth."
        ),
        thirty_second_pitch=(
            "ELIORA opens golden and sparkling, blooms into creamy white florals, then settles into musk, amber, "
            "and soft woods. It feels intimate, mysterious, and quietly exclusive."
        ),
        try_when="When someone asks for limited, special, evening, sensual, feminine, floral, musky, or warm.",
        trainer_note="Make it feel like a reveal. Do not over-explain it; invite them into the secret.",
    ),
}


def get_notes():
    return [
        builder.parse_notes(name, builder.SOURCE_FILES[name])
        for name in ["AUREYA", "ZEPHYR", "MARIS", "ELIORA"]
    ]


def brand_book_points() -> list[str]:
    fallback = [
        "TARA is a modern fragrance house built on mythology, identity, and sensory psychology.",
        "Each fragrance is designed as an extension of personal identity, not just a product.",
        "Luxury is subtle, felt, and remembered rather than announced.",
        "TARA builds identity through emotional connection, quiet confidence, and scent memory.",
    ]
    if not BRAND_BOOK.exists():
        return fallback

    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(BRAND_BOOK)).pages)
    wanted = [
        "TARA is a modern fragrance house built on the fusion of mythology, identity, and sensory psychology.",
        "Each fragrance is designed not as a product, but as an extension of personal identity.",
        "Luxury is subtle, felt, and remembered rather than announced.",
        "Each fragrance is designed to be worn quietly but remembered clearly.",
    ]
    return [line for line in wanted if line in text] or fallback


def event_rows() -> list[list[str]]:
    return [
        ["Field", "Details"],
        ["Crew", CREW_FULL_NAME],
        ["Role", CREW_ROLE],
        ["Trail stop", EVENT_NAME],
        ["Venue", EVENT_VENUE],
        ["Dates", EVENT_DATES],
        ["Hours", EVENT_HOURS],
        ["Status", EVENT_STATUS],
    ]


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ["top", "left", "bottom", "right"]:
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 130, bottom: int = 90, end: int = 130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    width_dxa = [int(width * 1440) for width in widths]
    total_dxa = sum(width_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total_dxa))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in width_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for col_index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width_dxa[col_index]))


def set_run(run, *, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def para(doc: Document, text: str = "", *, before: int = 0, after: int = 6, align=None, size: float = 10.6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.22
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        set_run(run, size=size, color=INK)
    return paragraph


def heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run(run, size={1: 17, 2: 12.5, 3: 11}[level], color=GOLD if level <= 2 else BLACK, bold=True)
    return paragraph


def bullet(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run(run, size=10.2, color=INK)


def number(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run(run, size=10.2, color=INK)


def add_docx_table(doc: Document, rows: list[list[str]], widths: list[float], *, header: bool = True, fill: str = LIGHT_GRAY):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, widths)
    for col_index, width in enumerate(widths):
        table.columns[col_index].width = Inches(width)

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if header and row_index == 0:
                set_cell_fill(cell, "111111")
            elif fill and row_index % 2 == 0:
                set_cell_fill(cell, fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run(
                run,
                size=8.7 if not (header and row_index == 0) else 8.5,
                color=WHITE if header and row_index == 0 else INK,
                bold=header and row_index == 0,
            )
    para(doc, "", after=3)


def add_callout_docx(doc: Document, text: str, label: str | None = None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_fill(cell, LIGHT_GOLD)
    set_cell_border(cell)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if label:
        label_run = paragraph.add_run(f"{label.upper()}: ")
        set_run(label_run, size=8.8, color=GOLD, bold=True)
    body_run = paragraph.add_run(text)
    set_run(body_run, size=10.1, color=INK)
    para(doc, "", after=2)


def set_docx_header(section, title: str):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO_DOCX.exists():
        paragraph.add_run().add_picture(str(LOGO_DOCX), width=Inches(1.1))
        paragraph.add_run("   ")
    title_run = paragraph.add_run(title)
    set_run(title_run, size=8.2, color=MUTED, bold=True)


def build_docx(notes):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.76)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.22
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style_name].font.name = "Calibri"

    set_docx_header(doc.sections[0], HANDBOOK_TITLE)
    para(doc, "", after=26)
    cover_logo = doc.add_paragraph()
    cover_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO_DOCX.exists():
        cover_logo.add_run().add_picture(str(LOGO_DOCX), width=Inches(4.0))
    para(doc, "", after=14)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(HANDBOOK_TITLE)
    set_run(title_run, size=25, color=BLACK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(HANDBOOK_SUBTITLE)
    set_run(subtitle_run, size=12.5, color=GOLD, bold=True)
    add_callout_docx(
        doc,
        "A friendly, practical guide for learning TARA, speaking about the scents with confidence, and helping customers find the fragrance that feels like them.",
        "Purpose",
    )
    add_docx_table(doc, event_rows(), [1.45, 5.05])
    para(doc, f"Prepared for {CREW_FIRST_NAME} | {DISPLAY_DATE}", align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_docx_header(doc.sections[-1], f"Start Here | {HANDBOOK_TITLE}")
    heading(doc, "Welcome To TARA", 1)
    add_callout_docx(doc, "TARA is worn quietly but remembered clearly.", "Brand line")
    heading(doc, "Assignment Snapshot", 2)
    add_docx_table(doc, event_rows(), [1.45, 5.05])
    heading(doc, "What TARA Stands For", 2)
    for point in brand_book_points():
        bullet(doc, point)
    heading(doc, "How To Speak TARA", 2)
    add_docx_table(
        doc,
        [
            ["Use", "Avoid"],
            ["Mood, memory, identity, skin feel, confidence, warmth, radiance.", "Clone, dupe, inspired-by, mixed from, formula ratios, or technical source claims."],
            ["Quiet luxury language: subtle, remembered, composed, luminous.", "Overpromising projection, medical claims, or comparing directly with competitor brands."],
            ["Ask what the customer wants to feel before recommending.", "Starting with a long ingredient list before understanding the customer."],
        ],
        [3.25, 3.25],
    )

    heading(doc, "Social Handles To Follow", 2)
    for item in [
        "Instagram: @tara_scents.my",
        "Website: tarascents.com",
        "Link hub: tarascents.com/links",
        "WhatsApp Concierge: +60 11-4304 2883",
        "Email: hello@tarascents.com",
    ]:
        bullet(doc, item)

    heading(doc, "Pop-Up Conversation Flow", 1)
    steps = [
        "Welcome warmly: ask if they prefer fresh, warm, floral, clean, subtle, or bold.",
        "Match mood first: recommend one or two scents based on what they want to feel.",
        "Spray and pause: let the opening settle before explaining the heart and drydown.",
        "Use a short story: one sentence is better than a lecture.",
        "Close gently: ask which scent felt most like them and offer the right size or set.",
    ]
    for step in steps:
        number(doc, step)
    add_callout_docx(
        doc,
        "Best question to ask: What do you want your scent to say before you say anything?",
        "Customer prompt",
    )

    heading(doc, "One-Glance Scent Wardrobe", 1)
    rows = [["Scent", "Role", "Best Customer Cue", "Short Direction"]]
    for note in notes:
        training = SCENT_TRAINING[note.name]
        title = note.name if note.name != "ELIORA" else "ELIORA - Special Launch"
        rows.append([title, training.identity, training.try_when, note.short_description])
    add_docx_table(doc, rows, [0.95, 1.6, 1.9, 2.05])

    for note in notes:
        title = note.name if note.name != "ELIORA" else "ELIORA - Special Launch"
        training = SCENT_TRAINING[note.name]
        doc.add_section(WD_SECTION.NEW_PAGE)
        set_docx_header(doc.sections[-1], f"{title} | {HANDBOOK_TITLE}")
        heading(doc, title, 1)
        add_callout_docx(doc, training.identity, "Identity")
        heading(doc, "30-Second Pitch", 2)
        para(doc, training.thirty_second_pitch)
        heading(doc, "Story Anchor", 2)
        para(doc, training.story_anchor)
        heading(doc, "Recommend When", 2)
        for item in training.customer_fit:
            bullet(doc, item)
        heading(doc, "Scent Journey", 2)
        script_rows = [["Moment", f"What {CREW_FIRST_NAME} Can Say"]]
        for cue_name in ["Opening", "Heart", "Drydown"]:
            if cue_name in note.script_cues:
                script_rows.append([cue_name, note.script_cues[cue_name]])
        add_docx_table(doc, script_rows, [1.1, 5.4])
        heading(doc, "Marketing Note Pyramid", 2)
        add_docx_table(doc, note.pyramid, [1.15, 5.35])
        heading(doc, "Trainer Note", 2)
        add_callout_docx(doc, training.trainer_note, "Remember")

    doc.add_section(WD_SECTION.NEW_PAGE)
    set_docx_header(doc.sections[-1], f"Practice Kit | {HANDBOOK_TITLE}")
    heading(doc, "Practice Kit", 1)
    heading(doc, "Quick Matching Prompts", 2)
    add_docx_table(
        doc,
        [
            ["Customer says", "Start with"],
            ["I want something soft and feminine.", "AUREYA"],
            ["I want clean, masculine, and confident.", "ZEPHYR"],
            ["I want unisex, calm, and not too loud.", "MARIS"],
            ["I want something special or evening-ready.", "ELIORA"],
        ],
        [3.4, 3.1],
    )
    heading(doc, "Mini Self-Check", 2)
    for item in [
        "Can I explain each scent in one sentence?",
        "Can I describe opening, heart, and drydown without reading?",
        "Can I ask a customer what they want to feel before recommending?",
        "Can I avoid clone, dupe, inspired-by, and formula language?",
        "Can I invite customers to follow @tara_scents.my before they leave?",
    ]:
        bullet(doc, item)

    footer = doc.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"TARA Scent Trail | {HANDBOOK_TITLE}")
    set_run(footer_run, size=8.3, color=MUTED)

    HANDBOOK_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)
    return DOCX_OUT


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CoverTitle",
            parent=styles["Title"],
            fontName="Times-Bold",
            fontSize=30,
            leading=34,
            textColor=PDF_BLACK,
            alignment=TA_CENTER,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CoverSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10.8,
            leading=15,
            textColor=PDF_GOLD,
            alignment=TA_CENTER,
            spaceAfter=17,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H1",
            parent=styles["Heading1"],
            fontName="Times-Bold",
            fontSize=21,
            leading=25,
            textColor=PDF_BLACK,
            spaceBefore=8,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H2",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=10.4,
            leading=13,
            textColor=PDF_GOLD,
            spaceBefore=10,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.2,
            leading=12.6,
            textColor=PDF_INK,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Small",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.1,
            leading=10.7,
            textColor=PDF_INK,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Callout",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=12.3,
            textColor=PDF_INK,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableText",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.2,
            leading=9.1,
            textColor=PDF_INK,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TableHeader",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7.2,
            leading=9.1,
            textColor=colors.white,
        )
    )
    return styles


def pp(text: str, style):
    safe = escape(str(text or "-")).replace("\n", "<br/>")
    return Paragraph(safe, style)


PAGE_WIDTH, PAGE_HEIGHT = letter


def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(colors.white)
    canvas.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, stroke=0, fill=1)
    canvas.setFillColor(PDF_BLACK)
    canvas.rect(0, PAGE_HEIGHT - 74, PAGE_WIDTH, 74, stroke=0, fill=1)
    if LOGO_PDF.exists():
        canvas.drawImage(str(LOGO_PDF), 44, PAGE_HEIGHT - 53, width=130, height=28.6, mask="auto")
    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.setFillColor(colors.white)
    canvas.drawRightString(PAGE_WIDTH - 44, PAGE_HEIGHT - 43, HANDBOOK_TITLE.upper())
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(PDF_MUTED)
    canvas.drawString(44, 26, f"TARA Scent Trail | Page {doc.page}")
    canvas.restoreState()


def callout_pdf(text: str, styles, label: str | None = None):
    label_text = f"{label.upper()}: " if label else ""
    table = Table(
        [[pp(f"{label_text}{text}", styles["Callout"])]],
        colWidths=[6.2 * inch],
        hAlign="CENTER",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), PDF_LIGHT_GOLD),
                ("BOX", (0, 0), (-1, -1), 0.55, PDF_LINE),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    return [table, Spacer(1, 6)]


def bullet_pdf(items: list[str], styles):
    return ListFlowable(
        [ListItem(pp(item, styles["Body"]), leftIndent=14) for item in items],
        bulletType="bullet",
        start="circle",
        leftIndent=18,
        bulletFontName="Helvetica",
        bulletFontSize=7,
        bulletColor=PDF_GOLD,
    )


def numbered_pdf(items: list[str], styles):
    return ListFlowable(
        [ListItem(pp(item, styles["Body"]), leftIndent=14) for item in items],
        bulletType="1",
        leftIndent=19,
        bulletFontName="Helvetica-Bold",
        bulletFontSize=7,
        bulletColor=PDF_GOLD,
    )


def table_pdf(rows: list[list[str]], widths: list[float], styles):
    styled_rows = []
    for row_index, row in enumerate(rows):
        styled_rows.append(
            [
                pp(cell, styles["TableHeader"] if row_index == 0 else styles["TableText"])
                for cell in row
            ]
        )
    table = Table(styled_rows, colWidths=[width * inch for width in widths], hAlign="CENTER", repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_BLACK),
                ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PDF_LIGHT_GRAY]),
                ("BOX", (0, 0), (-1, -1), 0.4, PDF_LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, PDF_LINE),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
            ]
        )
    )
    return [table, Spacer(1, 7)]


def add_pdf_start(story, notes, styles):
    story.append(Paragraph("Welcome To TARA", styles["H1"]))
    story.extend(callout_pdf("TARA is worn quietly but remembered clearly.", styles, "Brand line"))
    story.append(Paragraph("Assignment Snapshot", styles["H2"]))
    story.extend(table_pdf(event_rows(), [1.35, 4.85], styles))
    story.append(Paragraph("What TARA Stands For", styles["H2"]))
    story.append(bullet_pdf(brand_book_points(), styles))
    story.append(Spacer(1, 6))
    story.append(Paragraph("How To Speak TARA", styles["H2"]))
    story.extend(
        table_pdf(
            [
                ["Use", "Avoid"],
                ["Mood, memory, identity, skin feel, confidence, warmth, radiance.", "Clone, dupe, inspired-by, mixed from, formula ratios, or technical source claims."],
                ["Quiet luxury language: subtle, remembered, composed, luminous.", "Overpromising projection, medical claims, or direct competitor comparisons."],
                ["Ask what the customer wants to feel before recommending.", "Starting with a long ingredient list before understanding the customer."],
            ],
            [3.1, 3.1],
            styles,
        )
    )
    story.append(Paragraph("Social Handles To Follow", styles["H2"]))
    story.append(
        bullet_pdf(
            [
                "Instagram: @tara_scents.my",
                "Website: tarascents.com",
                "Link hub: tarascents.com/links",
                "WhatsApp Concierge: +60 11-4304 2883",
                "Email: hello@tarascents.com",
            ],
            styles,
        )
    )
    story.append(PageBreak())

    story.append(Paragraph("Pop-Up Conversation Flow", styles["H1"]))
    story.append(
        numbered_pdf(
            [
                "Welcome warmly: ask if they prefer fresh, warm, floral, clean, subtle, or bold.",
                "Match mood first: recommend one or two scents based on what they want to feel.",
                "Spray and pause: let the opening settle before explaining the heart and drydown.",
                "Use a short story: one sentence is better than a lecture.",
                "Close gently: ask which scent felt most like them and offer the right size or set.",
            ],
            styles,
        )
    )
    story.append(Spacer(1, 6))
    story.extend(
        callout_pdf(
            "Best question to ask: What do you want your scent to say before you say anything?",
            styles,
            "Customer prompt",
        )
    )
    story.append(Paragraph("One-Glance Scent Wardrobe", styles["H2"]))
    rows = [["Scent", "Role", "Best Customer Cue", "Short Direction"]]
    for note in notes:
        training = SCENT_TRAINING[note.name]
        title = note.name if note.name != "ELIORA" else "ELIORA - Special Launch"
        rows.append([title, training.identity, training.try_when, note.short_description])
    story.extend(table_pdf(rows, [0.78, 1.45, 1.8, 2.17], styles))


def add_pdf_scent(story, note, styles):
    title = note.name if note.name != "ELIORA" else "ELIORA - Special Launch"
    training = SCENT_TRAINING[note.name]
    story.append(PageBreak())
    story.append(Paragraph(title, styles["H1"]))
    story.extend(callout_pdf(training.identity, styles, "Identity"))
    story.append(Paragraph("30-Second Pitch", styles["H2"]))
    story.append(pp(training.thirty_second_pitch, styles["Body"]))
    story.append(Paragraph("Story Anchor", styles["H2"]))
    story.append(pp(training.story_anchor, styles["Body"]))
    story.append(Paragraph("Recommend When", styles["H2"]))
    story.append(bullet_pdf(training.customer_fit, styles))
    story.append(Spacer(1, 5))
    story.append(Paragraph("Scent Journey", styles["H2"]))
    script_rows = [["Moment", f"What {CREW_FIRST_NAME} Can Say"]]
    for cue_name in ["Opening", "Heart", "Drydown"]:
        if cue_name in note.script_cues:
            script_rows.append([cue_name, note.script_cues[cue_name]])
    story.extend(table_pdf(script_rows, [0.9, 5.3], styles))
    story.append(Paragraph("Marketing Note Pyramid", styles["H2"]))
    story.extend(table_pdf(note.pyramid, [1.0, 5.2], styles))
    story.extend(callout_pdf(training.trainer_note, styles, "Trainer note"))


def add_pdf_practice(story, styles):
    story.append(PageBreak())
    story.append(Paragraph("Practice Kit", styles["H1"]))
    story.append(Paragraph("Quick Matching Prompts", styles["H2"]))
    story.extend(
        table_pdf(
            [
                ["Customer says", "Start with"],
                ["I want something soft and feminine.", "AUREYA"],
                ["I want clean, masculine, and confident.", "ZEPHYR"],
                ["I want unisex, calm, and not too loud.", "MARIS"],
                ["I want something special or evening-ready.", "ELIORA"],
            ],
            [4.6, 1.6],
            styles,
        )
    )
    story.append(Paragraph("Mini Self-Check", styles["H2"]))
    story.append(
        bullet_pdf(
            [
                "Can I explain each scent in one sentence?",
                "Can I describe opening, heart, and drydown without reading?",
                "Can I ask a customer what they want to feel before recommending?",
                "Can I avoid clone, dupe, inspired-by, and formula language?",
                "Can I invite customers to follow @tara_scents.my before they leave?",
            ],
            styles,
        )
    )
    story.append(Spacer(1, 10))
    story.extend(
        callout_pdf(
            f"{CREW_FIRST_NAME} does not need to sound like a perfume encyclopedia. She just needs to be warm, curious, clear, and on-brand.",
            styles,
            "Final coaching note",
        )
    )


def build_pdf(notes):
    HANDBOOK_DIR.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=1.15 * inch,
        bottomMargin=0.58 * inch,
        title=HANDBOOK_TITLE,
        author="TARA",
    )
    story = []
    story.append(Spacer(1, 30))
    story.append(Paragraph(HANDBOOK_TITLE, styles["CoverTitle"]))
    story.append(Paragraph(HANDBOOK_SUBTITLE, styles["CoverSubtitle"]))
    story.extend(
        callout_pdf(
            "A friendly, practical guide for learning TARA, speaking about the scents with confidence, and helping customers find the fragrance that feels like them.",
            styles,
            "Purpose",
        )
    )
    story.extend(table_pdf(event_rows(), [1.35, 4.85], styles))
    story.append(Paragraph(f"Prepared for {CREW_FIRST_NAME} | {DISPLAY_DATE}", styles["CoverSubtitle"]))
    story.append(PageBreak())
    add_pdf_start(story, notes, styles)
    for note in notes:
        add_pdf_scent(story, note, styles)
    add_pdf_practice(story, styles)

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return PDF_OUT


def main():
    notes = get_notes()
    docx_path = build_docx(notes)
    pdf_path = build_pdf(notes)
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
